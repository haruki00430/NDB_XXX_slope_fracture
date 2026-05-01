from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def main() -> None:
    out_path = (
        Path(__file__).resolve().parents[2]
        / "04_Manuscripts"
        / "revision_comparison_table_20260430.xlsx"
    )

    wb = Workbook()
    ws_info = wb.active
    assert ws_info is not None
    ws_info.title = "提出情報"

    info_rows = [
        ("修正対照表（提出用・清書版対応）", ""),
        ("", ""),
        ("対象論文", "Terrain slope and fracture surgery rates in Japan (NDB ecological study)"),
        ("対象Word（清書版）", "Manuscript_slope_fracture_20260425_saito0430_清書版.docx"),
        ("対象原稿（編集元）", "Manuscript_slope_fracture.qmd"),
        ("参照ログ", "itemized_revision_log_branch2.md"),
        ("参照チェック", "itemized_revision_checklist_branch2.md"),
        ("作成日", str(date.today())),
        (
            "備考",
            "No.1-12の項目対応に、関連する文言微修正（No.8 wording、No.11 grammar、Results解釈文削除、Discussion断定緩和、Methods短文化）を統合して記載",
        ),
    ]

    for r, row in enumerate(info_rows, 1):
        ws_info.cell(r, 1, row[0])
        ws_info.cell(r, 2, row[1])

    ws_info.merge_cells("A1:B1")
    ws_info["A1"].font = Font(bold=True, size=13)
    ws_info["A1"].alignment = Alignment(horizontal="left")
    for c in ["A", "B"]:
        ws_info.column_dimensions[c].width = 48 if c == "B" else 24
    for r in range(1, len(info_rows) + 1):
        ws_info.cell(r, 1).alignment = Alignment(vertical="top", wrap_text=True)
        ws_info.cell(r, 2).alignment = Alignment(vertical="top", wrap_text=True)

    ws = wb.create_sheet("修正対照表")
    headers = [
        "番号",
        "指摘事項（要約）",
        "修正内容（反映事項）",
        "対応箇所（清書版Wordで確認する章・節・該当文）",
        "根拠ログ/コミット",
    ]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9E1F2")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    rows = [
        (
            "No.1",
            "高齢化率・歩行速度の位置づけ（交絡/媒介）",
            "aging rate / fast walking rate を formal mediator ではなく confounder-adjustment variables と明記。",
            "Methods > Statistical Analysis（Model 1/2説明直下）\n該当文: In the primary causal interpretation, aging rate and fast walking rate were treated as confounder-adjustment variables...",
            "log: No.1 / checklist: Done / commit: No.1 commit",
        ),
        (
            "No.2",
            "サンプルサイズとモデル安定性（N=47）",
            "小標本での係数安定性・多重共線性診断（condition number含む）への言及を追記。",
            "Methods > Statistical Analysis\n該当文: Given N = 47 with multiple covariates, we also checked coefficient stability and multicollinearity diagnostics...",
            "log: No.2 / checklist: Done / commit: No.2 commit",
        ),
        (
            "No.3",
            "アウトカムは発症ではなく手術率",
            "手術率アウトカムである点と、治療方針/アクセス差の影響可能性をLimitationsへ明示。",
            "Discussion > Limitations（Second）\n該当文: our outcome reflects surgery rates rather than true fracture incidence...",
            "log: No.3 / checklist: Done / commit: No.3 commit",
        ),
        (
            "No.4",
            "未調整交絡（気候・医療資源）",
            "未測定交絡に climate (snowfall and icing) を追加。",
            "Discussion > Limitations（Fourth）\n該当句: climate (snowfall and icing)",
            "log: No.4 / checklist: Done / commit: No.4 commit",
        ),
        (
            "No.5",
            "空間相関の可能性",
            "大平先生方針に従い、本文の追加変更は行わず運用記録で対応。",
            "本文追加なし（運用ファイルで対応）",
            "log: No.5 / checklist: Done / commit: 60df8a6",
        ),
        (
            "No.6",
            "線形性仮定",
            "大平先生方針に従い、本文新規追記は行わず既存記載で対応。",
            "Results > Regression Analysis: Hip Fracture（Residual diagnostics記載を維持）",
            "log: No.6 / checklist: Done / commit: ca55196",
        ),
        (
            "No.7",
            "曝露指標の粗さ",
            "既存Limitations記載（prefecture-level exposure, micro-scale未捕捉）を確認し非修正対応。",
            "Discussion > Limitations（Seventh）\n該当文: does not capture micro-scale terrain variations...",
            "log: No.7 / checklist: Done / commit: 555fb97",
        ),
        (
            "No.8",
            "年齢構成不一致（歩行速度40–74 vs 骨折高齢）",
            "Discussion/Limitationsで40–74歳公開データ制約を明示。加えて wording fix を実施（for which→because 含む）。",
            "Discussion > Walking Speed as a Mediator\nLimitations（Sixth）\n該当文: publicly available prefecture-level data were limited to individuals aged 40–74 years...",
            "log: No.8 + wording fixes / checklist: Done / commits: 6fad9ac, 668a335, 149791e",
        ),
        (
            "No.9 + No.12",
            "年齢標準化不足・高齢化率調整のみでは不十分",
            "マスキング補完シナリオ（\"-\"=0/5/9）で間接法SIR/ISR感度分析を追記。Methods/Results/Discussion/Limitationsへ統合。",
            "Methods > Statistical Analysis（SIR/ISR workflow）\nResults > Regression Analysis: Hip Fracture（0/5/9のβ/CI/p）\nDiscussion > Role of Aging Rate\nLimitations（Fifth）",
            "log: No.12(No.9同時) / checklist: No.9,12 Done / commit: 1f3c699",
        ),
        (
            "No.10",
            "結果解釈の強さ（因果を弱める）",
            "demonstrated→showed, amplifying→potentially amplifying, remained positively associated→was positively associated 等に調整。",
            "Discussion > Main Findings / Conclusions\n該当文の動詞トーンを慎重化",
            "log: No.10 / checklist: Done / commits: 93db0e1, 7145c84",
        ),
        (
            "No.11",
            "biological plausibility の断定トーン",
            "may/partly/suggested/potentially へ調整。文法微修正（which are common fall mechanisms）も反映。",
            "Discussion > Biological Plausibility\n該当文: ...may be biologically plausible... / which are common fall mechanisms...",
            "log: No.11 / checklist: Done / commits: cfa73dc, a20296d",
        ),
        (
            "追補A",
            "Resultsでの解釈文削除（IMRAD整合）",
            "Results内の interpret 文（We therefore interpret...）を削除。",
            "Results > Regression Analysis: Hip Fracture\nHC3/bootstrap記述直後",
            "log: Results文体微修正 / commit: d084cdf",
        ),
        (
            "追補B",
            "SKILL準拠強化（Discussion断定緩和・Methods短文化）",
            "Discussionの断定緩和を追加実施し、MethodsのSIR/ISR説明を短文化・分割。",
            "Discussion > Comparison/Role of Aging Rate\nMethods > Statistical Analysis（No.9/12追記段落）",
            "commit: 1d7ea7f",
        ),
    ]

    for r_idx, row in enumerate(rows, 2):
        for c_idx, value in enumerate(row, 1):
            cell = ws.cell(r_idx, c_idx, value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    widths = [10, 34, 52, 60, 44]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    thin = Side(style="thin", color="BFBFBF")
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            ws.cell(r, c).border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.freeze_panes = "A2"
    wb.save(out_path)
    print(str(out_path))


if __name__ == "__main__":
    main()
