from datetime import date
from pathlib import Path
from typing import cast

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(
    r"C:\Users\user\SharedWorkspace\projects\NDB_Research_Hub\projects\NDB_XXX_slope_fracture"
)
RESULTS = ROOT / "03_Analysis" / "results"
OUT_XLSX = ROOT / "04_Manuscripts" / "再解析・感度分析サマリ_査読対応_20260430.xlsx"


def style_header(ws, row_idx: int, max_col: int) -> None:
    fill = PatternFill("solid", fgColor="D9E1F2")
    for c in range(1, max_col + 1):
        cell = ws.cell(row_idx, c)
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_grid(ws) -> None:
    thin = Side(style="thin", color="BFBFBF")
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            ws.cell(r, c).border = Border(left=thin, right=thin, top=thin, bottom=thin)
            ws.cell(r, c).alignment = Alignment(vertical="top", wrap_text=True)


def write_df(ws, start_row: int, start_col: int, df: pd.DataFrame) -> int:
    for i, col in enumerate(df.columns, start_col):
        ws.cell(start_row, i, col)
    style_header(ws, start_row, start_col + len(df.columns) - 1)

    for r_off, row in enumerate(df.itertuples(index=False), 1):
        for c_off, val in enumerate(row, 0):
            ws.cell(start_row + r_off, start_col + c_off, val)
    return start_row + len(df) + 1


def main() -> None:
    sir = pd.read_csv(RESULTS / "sir_masking_sensitivity_summary.csv")
    vif = pd.read_csv(RESULTS / "multicollinearity_vif.csv")

    # Presented values used in manuscript (No.9/12 response)
    summary_rows = [
        {
            "区分": "主解析（Model 2, OLS）",
            "係数β（slope）": 3.4880,
            "95%CI下限": 0.112,
            "95%CI上限": 6.864,
            "p値": 0.043,
            "出典": "regression_results.txt",
        },
        {
            "区分": "主解析感度（HC3）",
            "係数β（slope）": 3.4880,
            "95%CI下限": -0.192,
            "95%CI上限": 7.168,
            "p値": 0.063,
            "出典": "regression_results.txt",
        },
        {
            "区分": "主解析感度（Bootstrap, B=5000）",
            "係数β（slope）": 3.4880,
            "95%CI下限": -0.03,
            "95%CI上限": 7.09,
            "p値": None,
            "出典": "Manuscript_slope_fracture.qmd 記載値",
        },
    ]
    integrated = pd.DataFrame(summary_rows)

    sir_view = sir.copy()
    sir_view["β（Model 2）"] = sir_view["beta_m2"].round(6)
    sir_view["95%CI（Model 2）"] = (
        sir_view["ci_low_m2"].round(6).astype(str)
        + " to "
        + sir_view["ci_high_m2"].round(6).astype(str)
    )
    sir_view["p（Model 2）"] = sir_view["p_m2"].round(6)
    sir_view["β（Model 1）"] = sir_view["beta_m1"].round(6)
    sir_view["95%CI（Model 1）"] = (
        sir_view["ci_low_m1"].round(6).astype(str)
        + " to "
        + sir_view["ci_high_m1"].round(6).astype(str)
    )
    sir_view["p（Model 1）"] = sir_view["p_m1"].round(6)
    sir_cols = [
        "scenario",
        "mask_fill",
        "age_groups_used",
        "nat_crude_per100k",
        "β（Model 2）",
        "95%CI（Model 2）",
        "p（Model 2）",
        "β（Model 1）",
        "95%CI（Model 1）",
        "p（Model 1）",
    ]
    sir_view = cast(pd.DataFrame, sir_view[sir_cols])

    wb = Workbook()

    # Sheet 1: 概要
    ws0 = wb.active
    if ws0 is None:
        raise RuntimeError("openpyxl Workbook has no active worksheet")
    ws0.title = "概要"
    info = [
        ("本ブックの目的", "寳澤先生コメント対応として実施した再解析・感度分析の内容と主要結果を一覧化"),
        ("作成日", str(date.today())),
        ("対象プロジェクト", "NDB_XXX_slope_fracture"),
        ("対象原稿", "Manuscript_slope_fracture.qmd"),
        ("主対象コメント", "No.9（年齢標準化不足）, No.12（高齢化率調整のみでは不十分）"),
        ("主解析の位置づけ", "主解析は据え置き。追加解析は感度分析（robustness check）"),
        ("SIR感度シナリオ", "A: '-'=0, B: '-'=5, C: '-'=9"),
        ("主要ファイル", "sir_masking_sensitivity_summary.csv, sir_masking_sensitivity_report.txt, age_standardized_indirect_regression.txt"),
    ]
    for i, (k, v) in enumerate(info, 1):
        ws0.cell(i, 1, k)
        ws0.cell(i, 2, v)
    ws0.merge_cells("A1:B1")
    ws0["A1"].font = Font(bold=True, size=13)
    ws0.column_dimensions["A"].width = 28
    ws0.column_dimensions["B"].width = 90
    style_grid(ws0)

    # Sheet 2: 再解析の全体像
    ws1 = wb.create_sheet("再解析の全体像")
    rows = [
        ("査読ID", "実施解析", "目的", "モデル/規模", "主結果（要約）", "出典"),
        (
            "No.9/12",
            "間接法SIR/ISR 感度分析（マスキング補完）",
            "年齢構成差による影響の頑健性確認",
            "A/B/Cの3シナリオ × Model1/Model2",
            "Model2のβは全シナリオで正（5.71, 3.61, 3.64）",
            "sir_masking_sensitivity_summary.csv",
        ),
        (
            "No.9/12",
            "主解析との比較（OLS/HC3/Bootstrap）",
            "推論の安定性確認",
            "N=47, Model2",
            "OLSでp=0.043、HC3とBootstrapではCIが0をまたぐ",
            "regression_results.txt / qmd",
        ),
        (
            "No.2連動",
            "多重共線性確認（VIF）",
            "推定不安定性の補足確認",
            "4説明変数",
            "全変数VIF<5（最大2.03）",
            "multicollinearity_vif.csv",
        ),
    ]
    for r, row in enumerate(rows, 1):
        for c, val in enumerate(row, 1):
            ws1.cell(r, c, val)
    style_header(ws1, 1, 6)
    for c, w in enumerate([12, 34, 34, 22, 44, 34], 1):
        ws1.column_dimensions[get_column_letter(c)].width = w
    ws1.freeze_panes = "A2"
    style_grid(ws1)

    # Sheet 3: 全再解析_統合一覧
    ws2 = wb.create_sheet("全再解析_統合一覧")
    ws2.cell(1, 1, "全再解析 統合一覧（主解析ロバスト性 + SIR感度分析）")
    ws2["A1"].font = Font(bold=True, size=12)
    ws2.merge_cells("A1:F1")
    next_row = write_df(ws2, 3, 1, integrated)
    write_df(ws2, next_row + 1, 1, sir_view)
    for c, w in enumerate([34, 20, 15, 15, 12, 34], 1):
        ws2.column_dimensions[get_column_letter(c)].width = w
    for c in range(7, 11):
        ws2.column_dimensions[get_column_letter(c)].width = 18
    ws2.freeze_panes = "A4"
    style_grid(ws2)

    # Sheet 4: SIR_シナリオ詳細
    ws3 = wb.create_sheet("SIR_シナリオ詳細")
    write_df(ws3, 1, 1, sir)
    ws3.cell(
        ws3.max_row + 2,
        1,
        "注: Public NDB masking includes possible secondary masking; A/B/C are pragmatic sensitivity bounds.",
    )
    for c, w in enumerate([14, 10, 44, 18, 12, 12, 12, 10, 12, 12, 12, 10], 1):
        ws3.column_dimensions[get_column_letter(c)].width = w
    ws3.freeze_panes = "A2"
    style_grid(ws3)

    # Sheet 5: 診断（VIF/補足）
    ws4 = wb.create_sheet("診断（VIF・補足）")
    ws4.cell(1, 1, "VIF 結果")
    ws4["A1"].font = Font(bold=True, size=12)
    write_df(ws4, 3, 1, vif)

    ws4.cell(10, 1, "補足（本文反映済みの要点）")
    ws4.cell(11, 1, "・No.9/12対応として、SIR/ISR感度分析（'-'=0/5/9）を Methods/Results/Discussion/Limitations に反映。")
    ws4.cell(12, 1, "・方向性（β符号）は安定だが、効果量は仮定に依存。")
    ws4.cell(13, 1, "・主解析は据え置きで、追加解析は頑健性確認として解釈。")
    for r in [10, 11, 12, 13]:
        ws4.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")
    for c, w in enumerate([24, 14, 12, 12], 1):
        ws4.column_dimensions[get_column_letter(c)].width = w
    style_grid(ws4)

    wb.save(OUT_XLSX)
    print(str(OUT_XLSX))


if __name__ == "__main__":
    main()
